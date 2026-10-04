"""Check the manuscript tables in `tables/` against the computed results.

Every number in Tables 1 and 3-7 (`tables/Table_*.docx`) is compared with the
value recomputed from the processed data (Table 1) or stored in `results/`
(Tables 3-7), after rounding to the precision printed in the table. The script
also checks the Model Confidence Set markers (†), the bold best-in-column
values, the leading models and the significance stars. Table 2 lists settings,
which are documented in the notebook's CONFIG cell.

Usage:
    python code/verify_tables.py
Exit status 0 means every checked cell agrees.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from descriptive_statistics import table1  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
TABLES = REPO / "tables"
RESULTS = REPO / "results"
HORIZONS = [1, 7, 14, 28]
MCS_ALPHA = 0.10

problems: list[str] = []
checked = 0


def check(label: str, printed: str, expected: str) -> None:
    global checked
    checked += 1
    if printed != expected:
        problems.append(f"{label}: table shows {printed!r}, results give {expected!r}")


def notebook_name(name: str) -> str:
    """Manuscript model name -> name used in the notebook and results/."""
    return (name.replace("rSTL-", "RobustSTL-").replace("N-BEATS", "NBEATS")
                .replace(" (direct)", "(direct)"))


def rows(table_no: int) -> list[list]:
    """Rows of a table as lists of (text, bold) cells, skipping header and group rows."""
    doc = Document(TABLES / f"Table_{table_no}.docx")
    out = []
    for row in doc.tables[0].rows[1:]:
        cells = [(c.text.strip(), any(r.bold for p in c.paragraphs for r in p.runs if r.text.strip()))
                 for c in row.cells]
        if len({t for t, _ in cells}) == 1:      # merged group-heading row
            continue
        out.append(cells)
    return out


def mcs_sets() -> dict[int, set[str]]:
    sets = {}
    for h in HORIZONS:
        mcs = pd.read_csv(RESULTS / f"MCS_h{h}.csv")
        sets[h] = set(mcs.loc[mcs["Pvalue"] >= MCS_ALPHA, "Model name"])
    return sets


def verify_table1() -> None:
    stats = table1()
    names = {"Close (USD)": "Close (USD)", "ln(Close)": "ln(Close)", "Log return": "Log return"}
    for cells in rows(1):
        series = names[cells[0][0]]
        s = stats.loc[series]
        if series == "Close (USD)":
            fmt = ["{:,.2f}"] * 4
        elif series == "ln(Close)":
            fmt = ["{:.3f}"] * 4
        else:
            fmt = ["{:.4f}", "{:.4f}", "{:.3f}", "{:.3f}"]
        cols = ["Mean", "SD", "Minimum", "Maximum"]
        for (text, _), col, f in zip(cells[1:5], cols, fmt):
            check(f"Table 1 {series} {col}", text, f.format(s[col]))
        check(f"Table 1 {series} Skewness", cells[5][0], f"{s['Skewness']:.3f}")
        check(f"Table 1 {series} Excess kurtosis", cells[6][0], f"{s['Excess kurtosis']:.3f}")
        check(f"Table 1 {series} ADF/KPSS", cells[7][0], f"{s['ADF p']:.3f} / {s['KPSS p']:.3f}")


def verify_metric_table(table_no: int, matrix: str, fmt: str, best: str) -> None:
    mat = pd.read_csv(RESULTS / matrix, index_col="Model")
    mat.columns = [int(c) for c in mat.columns]
    mcs = mcs_sets()
    table = rows(table_no)
    for cells in table:
        model = notebook_name(cells[0][0])
        for (text, _), h in zip(cells[1:], HORIZONS):
            value = mat.loc[model, h]
            dagger = "†" if model in mcs[h] else ""
            check(f"Table {table_no} {cells[0][0]} h={h}", text, fmt.format(value) + dagger)
    # bold marks the best value in each column (ties allowed at the stored precision)
    models = [notebook_name(c[0][0]) for c in table]
    for j, h in enumerate(HORIZONS, start=1):
        col = mat.loc[models, h].round(4)
        target = col.min() if best == "min" else col.max()
        for cells, model in zip(table, models):
            if cells[j][1]:
                check(f"Table {table_no} bold h={h}", f"{cells[0][0]}={col[model]}", f"{cells[0][0]}={target}")


def verify_table6() -> None:
    best = pd.read_csv(RESULTS / "best_per_horizon.csv").set_index("h")
    for cells in rows(6):
        h = int(cells[0][0])
        b = best.loc[h]
        check(f"Table 6 h={h} leader", notebook_name(cells[1][0]), b["Leader"])
        mean, sd = (int(v) for v in b["RMSE"].split("±"))
        check(f"Table 6 h={h} RMSE", cells[2][0], f"{mean:,} ± {sd:,}")
        check(f"Table 6 h={h} R2", cells[3][0], f"{b['R2']:.4f}")
        check(f"Table 6 h={h} MAPE", cells[4][0], f"{b['MAPE']:.2f}")
        check(f"Table 6 h={h} DA", cells[5][0], f"{b['DA']:.1f}")
        check(f"Table 6 h={h} in MCS", cells[6][0], "Yes" if b["in_MCS"] == "yes" else "No")


def stars(p: float) -> str:
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


def verify_table7() -> None:
    dm = pd.read_csv(RESULTS / "DM_seedavg.csv").set_index(["Baseline", "h"])
    for cells in rows(7):
        base = notebook_name(cells[0][0])
        for (text, _), h in zip(cells[1:], HORIZONS):
            r = dm.loc[(base, h)]
            check(f"Table 7 {cells[0][0]} h={h}", text, f"{r['DM']:.2f}{stars(r['p'])}")


def main() -> None:
    verify_table1()
    verify_metric_table(3, "RMSE_mean_matrix.csv", "{:,.0f}", "min")
    verify_metric_table(4, "R2_mean_matrix.csv", "{:.4f}", "max")
    verify_metric_table(5, "MAPE_mean_matrix.csv", "{:.2f}", "min")
    verify_table6()
    verify_table7()
    if problems:
        print(f"{len(problems)} of {checked} checks FAILED:")
        for p in problems:
            print("  -", p)
        sys.exit(1)
    print(f"All {checked} checks passed (Tables 1 and 3-7 agree with the computed results).")


if __name__ == "__main__":
    main()
